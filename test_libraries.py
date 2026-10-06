"""The font and figure libraries  (run: python test_libraries.py).

Bug hunt item 6. The two libraries every style, cover, wrap design and
manuscript draws on, and the routes that serve files. Checked, against
throwaway folders:

    UPLOADS   not a font / not an image, a broken file under a name a good one
              has, a different file under a name in use, the same file again,
              a built-in's name, too many pixels, a name put there by hand;
              nothing left behind in the folder
    DELETE    a font a style, a cover template and a wrap design name, a figure
              a book and a style's chapter art name: the page says where it is
              used and asks; unused, unknown and built-in
    CHECKS    what went missing is said: a book's figures, a cover template's
              fonts, a wrap design's text and pictures (build, package, and the
              designer page in a browser)
    SERVING   /out, /download and /figures/file refuse anything outside their folder
"""

import sys, os, io, json, re, shutil, tempfile, threading, subprocess, time

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass
import logging; logging.disable(logging.INFO)

import app as A
import engine, epub, manuscript, matter, wrap_design
from test_rich_editor import find_browser

fails = []
def check(name, cond, detail=''):
    print(('  ok   ' if cond else '  FAIL ') + name + (('  ' + str(detail)) if not cond else ''))
    if not cond:
        fails.append(name)


tmp = tempfile.mkdtemp(prefix='ts-libraries-')
ROOT = os.path.join(tmp, 'data')
os.makedirs(ROOT)
shutil.copytree(os.path.join(HERE, 'presets'), os.path.join(ROOT, 'presets'))
shutil.copytree(os.path.join(HERE, 'covers'), os.path.join(ROOT, 'covers'),
                ignore=shutil.ignore_patterns('assets'))
shutil.copytree(os.path.join(HERE, 'fonts'), os.path.join(ROOT, 'fonts'),
                ignore=shutil.ignore_patterns('licenses'))
A.PRESET_DIR, A.COVER_DIR = os.path.join(ROOT, 'presets'), os.path.join(ROOT, 'covers')
A.FONT_DIR = engine.FONT_DIR = wrap_design.FONT_DIR = os.path.join(ROOT, 'fonts')
A.FIGURE_DIR = engine.FIGURE_DIR = epub.FIGURE_DIR = manuscript.FIGURE_DIR = \
    os.path.join(ROOT, 'figures')
A.PROJECT_DIR = os.path.join(ROOT, 'projects')
A.PROJECT_MS_DIR = os.path.join(A.PROJECT_DIR, 'manuscripts')
A.HISTORY_DIR = os.path.join(A.PROJECT_DIR, 'history')
A.OUT_DIR = os.path.join(ROOT, 'out')
A.UPLOAD_DIR = os.path.join(ROOT, 'uploads')
A.COVER_THUMB_DIR = os.path.join(A.OUT_DIR, '_cover_thumbs')
A.PROJECT_THUMB_DIR = os.path.join(A.OUT_DIR, '_project_thumbs')
A.WRAP_DESIGN_DIR = os.path.join(A.OUT_DIR, '_wrap_designer')
A.COVER_ASSET_DIR = engine.COVER_ASSET_DIR = os.path.join(ROOT, 'covers', 'assets')
wrap_design.ART_DIRS = [A.COVER_ASSET_DIR, A.WRAP_DESIGN_DIR]
for d in (A.FIGURE_DIR, A.PROJECT_MS_DIR, A.HISTORY_DIR, A.UPLOAD_DIR, A.COVER_THUMB_DIR,
          A.PROJECT_THUMB_DIR, A.WRAP_DESIGN_DIR, A.COVER_ASSET_DIR):
    os.makedirs(d, exist_ok=True)

FONT = lambda n: content(os.path.join(HERE, 'fonts', n))


def png(w=40, h=30, colour=(40, 60, 90), mode='RGB'):
    from PIL import Image
    b = io.BytesIO()
    Image.new(mode, (w, h), colour if mode == 'RGB' else 0).save(b, 'PNG')
    return b.getvalue()


def book(pid, text, **extra):
    with open(os.path.join(A.PROJECT_MS_DIR, pid + '.md'), 'w', encoding='utf-8') as f:
        f.write(text)
    A.save_project_file(pid, {
        'name': pid, 'preset': 'classic-literary', 'title': 'The Salt Road',
        'author': 'Ellinor Vale', 'front_matter': 'none', 'right_hand_starts': False,
        'cover_mode': 'none', 'format': 'pdf', 'include_toc': False, 'smartquotes': True,
        **matter.blank(), 'manuscript_file': pid + '.md', 'manuscript_type': 'markdown',
        'print_retailer': 'kdp', 'print_binding': 'paperback', 'print_paper': 'cream',
        'last_pdf': '', 'last_epub': '', **extra})


TITLE = {'id': 'title', 'type': 'text', 'anchor': 'front', 'x': 0.5, 'y': 0.9, 'w': 5,
         'text': 'The Salt Road', 'font': 'Gone-Bold.ttf', 'size': 40, 'leading': 1.1,
         'color': '#fbf3e2', 'align': 'center'}
DESIGN = {'elements': [
    {'id': 'bg', 'type': 'rect', 'fill': 'back', 'color': '#20283b'},
    TITLE,
    {'id': 'author', 'type': 'text', 'anchor': 'front', 'x': 0.5, 'y': 7.5, 'w': 5,
     'text': 'Ellinor Vale', 'font': 'Shared-Regular.ttf', 'size': 18, 'leading': 1.1,
     'color': '#fbf3e2', 'align': 'center'},
    {'id': 'pic', 'type': 'image', 'anchor': 'back', 'x': 1, 'y': 1, 'w': 1, 'h': 1,
     'src': 'vanished.png'},
]}

# The browser half's harness: Flask takes no routes after its first request.
from flask import request
RESULT = {}
SCRIPT = r'''
var f = document.getElementById('f');
f.onload = function () {
  var w = f.contentWindow, d = f.contentDocument, out = {};
  var wait = setInterval(function () {
    if (!w.WD_READY) return;
    clearInterval(wait);
    out.checks = d.getElementById('wd-checks').textContent;
    w.WD_EDITOR.select(['title']);
    var sel = Array.prototype.filter.call(d.querySelectorAll('select'), function (s) {
      return Array.prototype.some.call(s.options, function (o) { return o.value === 'Gone-Bold.ttf'; });
    })[0];
    out.select_value = sel ? sel.value : null;
    out.select_label = sel ? sel.options[sel.selectedIndex].textContent : null;
    out.font_kept = w.WD_EDITOR.design().elements.find(function (e) { return e.id === 'title'; }).font;
    var x = new XMLHttpRequest(); x.open('POST', '/_lib_result', false); x.send(JSON.stringify(out));
  }, 100);
};'''

@A.app.route('/_lib_harness')
def _lib_harness():
    return ('<!doctype html><meta charset="utf-8"><iframe id="f" src="/wrap-designer?project=wrapbook"'
            ' style="width:1300px;height:900px"></iframe><script>' + SCRIPT + '</script>')

@A.app.route('/_lib_result', methods=['POST'])
def _lib_result():
    RESULT.update(json.loads(request.get_data(as_text=True)))
    return 'ok'


C = A.app.test_client()


def page_text(r):
    return re.sub(r'\s+', ' ', r.get_data(as_text=True))


def upload(kind, files, follow=True):
    return C.post(f'/{kind}/upload', data={kind: [(io.BytesIO(b), n) for n, b in files]},
                  content_type='multipart/form-data', follow_redirects=follow)


def content(path):
    try:
        with open(path, 'rb') as f:
            return f.read()
    except OSError:
        return None


def no_leftovers(folder):
    return [f for f in os.listdir(folder) if '.upload' in f or f.endswith('.tmp')]


def uploads_tests():
    print('\n[font uploads]')
    fd = A.FONT_DIR
    r = upload('fonts', [('Notafont.ttf', b'not a font' * 100), ('readme.txt', b'x')])
    check('not a font, and not a font file, are refused',
          not os.path.exists(os.path.join(fd, 'Notafont.ttf'))
          and 'Notafont.ttf (not an embeddable font' in page_text(r)
          and 'readme.txt (not a .ttf/.otf)' in page_text(r))
    upload('fonts', [('Shared-Regular.ttf', FONT('Lora-Regular.ttf'))])
    good = content(os.path.join(fd, 'Shared-Regular.ttf'))
    r = upload('fonts', [('Shared-Regular.ttf', b'broken' * 500)])
    check('a broken upload leaves the good font of that name alone',
          content(os.path.join(fd, 'Shared-Regular.ttf')) == good
          and not os.path.exists(os.path.join(fd, 'Shared-Regular-2.ttf')))
    r = upload('fonts', [('Shared-Regular.ttf', FONT('Vollkorn-Regular.ttf'))])
    check('a different font under that name goes in beside it, and says so',
          content(os.path.join(fd, 'Shared-Regular.ttf')) == good
          and content(os.path.join(fd, 'Shared-Regular-2.ttf'))
          == FONT('Vollkorn-Regular.ttf')
          and 'so this one is Shared-Regular-2.ttf' in page_text(r))
    before = sorted(os.listdir(fd))
    r = upload('fonts', [('Shared-Regular.ttf', FONT('Lora-Regular.ttf'))])
    check('the same font again is the one already there',
          sorted(os.listdir(fd)) == before and 'already in the library' in page_text(r))
    upload('fonts', [('Book-Regular.ttf', FONT('Lora-Regular.ttf'))])
    check('a built-in name taken by another font: the built-in is untouched',
          content(os.path.join(fd, 'Book-Regular.ttf')) == FONT('Book-Regular.ttf')
          and os.path.exists(os.path.join(fd, 'Book-Regular-2.ttf')))
    if os.path.exists(os.path.join(fd, 'Book-Regular-2.ttf')):
        os.remove(os.path.join(fd, 'Book-Regular-2.ttf'))
    check('nothing half-stored is left in the font folder', not no_leftovers(fd),
          no_leftovers(fd))

    print('\n[figure uploads]')
    gd = A.FIGURE_DIR
    upload('figures', [('map.png', png())])
    good = content(os.path.join(gd, 'map.png'))
    r = upload('figures', [('map.png', b'not an image')])
    check('a broken upload leaves the good figure of that name alone',
          content(os.path.join(gd, 'map.png')) == good
          and 'not a readable image' in page_text(r))
    r = C.post('/figures/upload', data={'figures': [(io.BytesIO(png(colour=(200, 0, 0))), 'map.png')]},
               content_type='multipart/form-data', headers={'X-Requested-With': 'fetch'})
    j = r.get_json()
    check('a different image under that name goes in beside it; the editor is told its name',
          j['saved'] == ['map-2.png'] and content(os.path.join(gd, 'map.png')) == good,
          j)
    r = C.post('/figures/upload', data={'figures': [(io.BytesIO(png()), 'map.png')]},
               content_type='multipart/form-data', headers={'X-Requested-With': 'fetch'})
    check('the same image again is the one already there', r.get_json()['saved'] == ['map.png'],
          r.get_json())
    r = upload('figures', [('huge.png', png(20000, 10000, mode='1'))])
    check('too many pixels is refused, and says why',
          not os.path.exists(os.path.join(gd, 'huge.png')) and 'too many pixels' in page_text(r))
    check('nothing half-stored is left in the figure folder', not no_leftovers(gd),
          no_leftovers(gd))

    print('\n[a figure named by hand]')
    with open(os.path.join(gd, 'My Map.png'), 'wb') as f:
        f.write(png(colour=(0, 200, 0)))
    r = C.get('/figures/file/My Map.png')
    check('it is shown', r.status_code == 200 and r.data == content(os.path.join(gd, 'My Map.png')))
    r.close()
    C.post('/figures/delete/My Map.png')
    check('and it can be removed', not os.path.exists(os.path.join(gd, 'My Map.png')))


def delete_tests():
    print('\n[deleting a font in use]')
    fd = A.FONT_DIR
    # a style, a cover template and a book's wrap design all name Shared-Regular.ttf
    st = A.load_preset('classic-literary')
    st['name'] = 'Classic “quoted”'
    st['font_files'] = dict(st['font_files'], regular='Shared-Regular.ttf')
    with open(os.path.join(A.PRESET_DIR, 'classic-literary.json'), 'w', encoding='utf-8') as f:
        json.dump(st, f)
    with open(os.path.join(A.COVER_DIR, 'fantasy-emerald.json'), encoding='utf-8') as f:
        tpl = json.load(f)
    tpl['fonts']['serif'] = 'Shared-Regular.ttf'
    with open(os.path.join(A.COVER_DIR, 'fantasy-emerald.json'), 'w', encoding='utf-8') as f:
        json.dump(tpl, f)
    book('wrapbook', '# One\n\nText.\n', cover_mode='wrap', wrap_design=DESIGN)
    r = C.post('/fonts/delete/Shared-Regular.ttf')
    t = page_text(r)
    check('the page says where it is used and keeps it',
          os.path.exists(os.path.join(fd, 'Shared-Regular.ttf')) and 'is in use' in t
          and 'the style “Classic “quoted””' in t
          and 'the cover template “' in t and 'the wrap design of “wrapbook”' in t, t[:300])
    C.post('/fonts/delete/Shared-Regular.ttf', data={'confirm': '1'})
    check('asked again, it goes', not os.path.exists(os.path.join(fd, 'Shared-Regular.ttf')))
    C.post('/fonts/delete/Shared-Regular-2.ttf')
    check('an unused font goes at once', not os.path.exists(os.path.join(fd, 'Shared-Regular-2.ttf')))
    r = C.post('/fonts/delete/Book-Regular.ttf', follow_redirects=True)
    check('a built-in stays', os.path.exists(os.path.join(fd, 'Book-Regular.ttf'))
          and 'ships with the app' in page_text(r))
    r = C.post('/fonts/delete/..%2Fpresets%2Fclassic-literary.json', follow_redirects=True)
    check('a name not in the library removes nothing',
          os.path.exists(os.path.join(A.PRESET_DIR, 'classic-literary.json'))
          and 'isn’t in the library' in page_text(r))
    shutil.copy(os.path.join(fd, 'Lora-Bold.ttf'), os.path.join(fd, "O'Brien Sans.ttf"))
    t = page_text(C.get('/fonts'))
    check('the remove prompt survives a name with a quote in it (put there by hand)',
          'data-ask="Remove O&#39;Brien Sans.ttf?' in t and "confirm('Remove" not in t,
          re.findall(r'.{60}Brien.{60}', t)[:2])
    C.post("/fonts/delete/O'Brien Sans.ttf")
    check('and it can be removed', not os.path.exists(os.path.join(fd, "O'Brien Sans.ttf")))

    print('\n[deleting a figure in use]')
    gd = A.FIGURE_DIR
    book('figbook', '# One\n\n~~~ figure src="map.png"\nThe map.\n~~~\n\nText.\n')
    st['chapter_art'] = dict(st.get('chapter_art') or {}, image='map-2.png')
    with open(os.path.join(A.PRESET_DIR, 'classic-literary.json'), 'w', encoding='utf-8') as f:
        json.dump(st, f)
    t = page_text(C.post('/figures/delete/map.png'))
    check('a book’s figure: the page says where, and keeps it',
          os.path.exists(os.path.join(gd, 'map.png')) and 'the book “figbook”' in t, t[:300])
    t = page_text(C.post('/figures/delete/map-2.png'))
    check('a style’s chapter art: the same',
          os.path.exists(os.path.join(gd, 'map-2.png')) and '(chapter art)' in t, t[:300])
    # a Word book's pictures go into the library under names of their own
    import docx
    d = docx.Document()
    d.add_heading('One', 1)
    d.add_paragraph('Text.')
    d.add_picture(io.BytesIO(png(colour=(9, 9, 200))))
    d.save(os.path.join(A.PROJECT_MS_DIR, 'wordbook.docx'))
    book('wordbook', '', manuscript_file='wordbook.docx', manuscript_type='file')
    A._project_manuscript_text(A.load_project('wordbook'), report_import=False)
    pic = next(f for f in os.listdir(gd) if f.startswith('wordbook-'))
    t = page_text(C.post(f'/figures/delete/{pic}'))
    check('a Word book’s picture: the same', 'the book “wordbook”' in t, t[:300])
    C.post('/figures/delete/map.png', data={'confirm': '1'})
    check('asked again, it goes', not os.path.exists(os.path.join(gd, 'map.png')))


def checks_tests():
    print('\n[what went missing, in the checks]')
    # figbook's map.png is gone now, and so is the cover template's Shared-Regular.ttf
    A.update_project('figbook', {'cover_mode': 'designed', 'cover_template': 'fantasy-emerald'})
    t = page_text(C.post('/project/figbook/generate'))
    check('a book build names its missing figure', 'box: map.png' in t,
          re.findall(r'Not in the figure.{0,160}', t)[:1])
    check('and its cover template’s missing font', 'Cover fonts' in t
          and 'Shared-Regular.ttf' in t, re.findall(r'Cover fonts.{0,160}', t)[:1])
    t = page_text(C.post('/project/figbook/print-package', data={'scope': 'print'}))
    check('so does its print package (the interior and the wrap)',
          'map.png' in t and 'Cover fonts' in t, t[:200])
    C.post('/project/wrapbook/generate')
    t = page_text(C.post('/project/wrapbook/print-package', data={'scope': 'print'}))
    check('a wrap design: text whose font is gone',
          'Text left off the cover' in t and 'Gone-Bold.ttf' in t and 'Shared-Regular.ttf' in t,
          re.findall(r'Text left.{0,160}', t)[:1])
    check('and a picture whose file is gone', 'Pictures left off the cover' in t
          and 'vanished.png' in t)
    good = engine.build_pdf(manuscript.parse_markdown('# One\n\nText.\n'),
                            A.load_preset('modern-clean'), os.path.join(tmp, 'ok.pdf'),
                            {'title': 'T', 'author': 'A', 'front_matter': 'none',
                             'cover_mode': 'none', **matter.blank()})
    check('a book with nothing missing reports nothing',
          good['figures_missing'] == [] and good['cover_fonts_missing'] == []
          and not A._missing_checks(good['figures_missing'], good['cover_fonts_missing']))


def serving_tests():
    print('\n[serving files]')
    with open(os.path.join(A.OUT_DIR, 'book.pdf'), 'wb') as f:
        f.write(b'%PDF-1.4 x')
    r = C.get('/download/book.pdf')
    check('a file in the folder is served', r.status_code == 200)
    r.close()
    bad = []
    for url in ('/download/../presets/classic-literary.json',
                '/download/..%2Fpresets%2Fclassic-literary.json',
                '/out/..%2F..%2Fdata%2Fpresets%2Fclassic-literary.json',
                '/out/%2E%2E/presets/classic-literary.json',
                '/download/' + os.path.join(A.PRESET_DIR, 'classic-literary.json').replace('\\', '/'),
                '/figures/file/..%2Fpresets%2Fclassic-literary.json',
                '/figures/file/../fonts/Book-Regular.ttf',
                '/wrap-designer/font/..%2Fpresets%2Fclassic-literary.json'):
        r = C.get(url)
        if r.status_code == 200:
            bad.append(url)
        r.close()
    check('nothing outside its folder is', not bad, bad)


def browser_tests():
    print('\n[the Wrap designer, in a browser]')
    browser = find_browser()
    if not browser:
        print('  SKIP no Chrome/Edge found (set TS_BROWSER)')
        return
    from werkzeug.serving import make_server
    server = make_server('127.0.0.1', 0, A.app, threaded=True)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    try:
        proc = subprocess.Popen([browser, '--headless=new', '--disable-gpu', '--no-first-run',
                                 '--user-data-dir=' + os.path.join(tmp, 'browser'),
                                 '--window-size=1400,1000',
                                 f'http://127.0.0.1:{server.server_port}/_lib_harness'],
                                stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        try:
            for _ in range(240):
                if RESULT:
                    break
                time.sleep(0.5)
        finally:
            proc.kill()
            proc.wait()
    finally:
        server.shutdown()
    R = RESULT
    check('the designer opened the book', bool(R), R)
    check('its checks say the text would be left off the cover',
          'left off the cover' in (R.get('checks') or '') and 'Gone-Bold.ttf' in R.get('checks', ''),
          R.get('checks'))
    check('the font picker shows the missing font, marked, not the first in the list',
          R.get('select_value') == 'Gone-Bold.ttf'
          and 'not in the font library' in (R.get('select_label') or ''), R)
    check('and the design still names it', R.get('font_kept') == 'Gone-Bold.ttf')


def main():
    try:
        uploads_tests()
        delete_tests()
        checks_tests()
        serving_tests()
        browser_tests()
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    return 0 if not fails else 1


if __name__ == '__main__':
    code = main()
    print('\n' + ('ALL PASS' if not fails else 'FAILED: ' + ', '.join(fails)))
    sys.exit(code)
