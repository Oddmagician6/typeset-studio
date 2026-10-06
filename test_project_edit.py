"""Book Edit page tests  (run: python test_project_edit.py).

Bug hunt item 3. The Edit page is the biggest form in the app and every book
goes through it, and it saves the book from the fields the page offered - so
whatever the page can't show, a save can quietly change. Checked:

    UNCHANGED  open the page and save it as a browser would, untouched: the book
               is the same, for books in every cover mode, with print settings,
               matter and uploads set, and with what the page has no option for
               (a style or cover template since deleted, a printer or paper this
               version doesn't know)
    UPLOADS    replacing the manuscript (kept in History first; a file that isn't
               a manuscript refused; text in other encodings read as written),
               the cover and the back-cover image (wrong kinds refused, and said
               so), clearing the back image
    VALUES     numbers that aren't numbers, out of range, or not finite
"""

import sys, os, json, time, shutil, tempfile, threading, subprocess, io
from html.parser import HTMLParser

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass
import logging; logging.disable(logging.INFO)

import app as A
import manuscript

fails = []
def check(name, cond, detail=''):
    print(('  ok   ' if cond else '  FAIL ') + name + (('  ' + str(detail)) if not cond else ''))
    if not cond:
        fails.append(name)


class FormFields(HTMLParser):
    """What a browser would submit from the form posting to `action`, untouched:
    each field's value as the page set it, ticked boxes only, the selected option
    (or the first, as a browser picks), no files. (As in test_upgrade.py.)"""
    def __init__(self, action):
        super().__init__(convert_charrefs=True)
        self.action, self.inside, self.fields = action, False, []
        self.select, self.options, self.textarea = None, [], None

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if tag == 'form':
            self.inside = self.action in (a.get('action') or '')
        if not self.inside or 'disabled' in a:
            return
        if tag == 'input' and a.get('name'):
            kind = (a.get('type') or 'text').lower()
            if kind in ('checkbox', 'radio'):
                if 'checked' in a:
                    self.fields.append((a['name'], a.get('value', 'on')))
            elif kind not in ('file', 'submit', 'button', 'image', 'reset'):
                self.fields.append((a['name'], a.get('value', '')))
        elif tag == 'select' and a.get('name'):
            self.select, self.options = a['name'], []
        elif tag == 'option' and self.select:
            self.options.append([a.get('value'), 'selected' in a, ''])
        elif tag == 'textarea' and a.get('name'):
            self.textarea = [a['name'], '']

    def handle_data(self, data):
        if self.textarea is not None:
            self.textarea[1] += data
        elif self.select and self.options and self.options[-1][0] is None:
            self.options[-1][2] += data

    def handle_endtag(self, tag):
        if tag == 'form':
            self.inside = False
        elif tag == 'textarea' and self.textarea is not None:
            name, text = self.textarea
            self.fields.append((name, text[1:] if text.startswith('\n') else text))
            self.textarea = None
        elif tag == 'select' and self.select:
            opts = [(v if v is not None else t.strip(), s) for v, s, t in self.options]
            pick = next((v for v, s in opts if s), opts[0][0] if opts else None)
            if pick is not None:
                self.fields.append((self.select, pick))
            self.select = None


def form_of(html, action):
    from werkzeug.datastructures import MultiDict
    p = FormFields(action)
    p.feed(html)
    return MultiDict(p.fields)


def use_folders(tmp):
    shutil.copytree(os.path.join(HERE, 'presets'), os.path.join(tmp, 'presets'))
    shutil.copytree(os.path.join(HERE, 'covers'), os.path.join(tmp, 'covers'),
                    ignore=shutil.ignore_patterns('assets'))
    A.PRESET_DIR, A.COVER_DIR = os.path.join(tmp, 'presets'), os.path.join(tmp, 'covers')
    A.PROJECT_DIR = os.path.join(tmp, 'projects')
    A.PROJECT_MS_DIR = os.path.join(A.PROJECT_DIR, 'manuscripts')
    A.HISTORY_DIR = os.path.join(A.PROJECT_DIR, 'history')
    A.OUT_DIR = os.path.join(tmp, 'out')
    A.UPLOAD_DIR = os.path.join(tmp, 'uploads')
    A.COVER_THUMB_DIR = os.path.join(A.OUT_DIR, '_cover_thumbs')
    A.PROJECT_THUMB_DIR = os.path.join(A.OUT_DIR, '_project_thumbs')
    for d in (A.PROJECT_MS_DIR, A.HISTORY_DIR, A.UPLOAD_DIR, A.COVER_THUMB_DIR,
              A.PROJECT_THUMB_DIR):
        os.makedirs(d, exist_ok=True)


def png_bytes(w=60, h=90, colour=(120, 30, 30)):
    from PIL import Image
    b = io.BytesIO()
    Image.new('RGB', (w, h), colour).save(b, 'PNG')
    return b.getvalue()


AWKWARD = 'A “quoted” line</textarea><b>x</b> {{ y }} & <\nsecond line\n\n third'

# Every field the page shows, set to something other than its default.
FULL = {
    'name': 'Full Book', 'preset': 'classic-literary', 'format': 'both',
    'include_toc': True, 'smartquotes': False, 'press': True,
    'title': 'The Salt Road', 'subtitle': 'A <Novel>', 'author': 'Ellinor Vale',
    'year': '2026', 'publisher': 'Ashforge & Sons', 'front_matter': 'title',
    'right_hand_starts': False, 'cover_overlay': True, 'cover_color': 'dark',
    'cover_mode': 'designed', 'cover_template': 'literary-ivory',
    'cover_collection': 'The Salt Cycle', 'cover_kicker': 'A tale',
    'cover_accent': 'E. Vale', 'cover_epigraph': 'Quote "here"', 'cover_studio': 'Ashforge',
    'print_retailer': 'ingramspark', 'print_binding': 'jacket', 'print_paper': 'cream',
    'print_blurb': 'Para one.\n\nPara two.', 'print_flap_blurb': 'Flap.',
    'print_flap_bio': 'Bio.', 'print_back_w': 2.25, 'print_back_y': 0.7,
}


def make_book(pid, **over):
    ms = pid + '.md'
    with open(os.path.join(A.PROJECT_MS_DIR, ms), 'w', encoding='utf-8') as f:
        f.write('# One\n\nThe text.\n')
    data = dict(FULL, manuscript_file=ms, manuscript_type='markdown', cover_file='',
                last_pdf='', last_epub='', created='2026-01-01T00:00:00',
                updated='2026-01-01T00:00:00', print_back_file='',
                **{k: f'{k} text' for k in A.matter.KEYS})
    data.update(over)
    A.save_project_file(pid, data)
    return data


def stored(pid):
    with open(os.path.join(A.PROJECT_DIR, A.secure_filename(pid) + '.json'),
              encoding='utf-8') as f:
        return json.load(f)


def unchanged_save(c, pid):
    html = c.get(f'/project/{pid}/edit').get_data(as_text=True)
    before = stored(pid)
    r = c.post(f'/project/{pid}/edit', data=form_of(html, f'/project/{pid}/edit'))
    after = stored(pid)
    diff = {k: (before.get(k), after.get(k)) for k in set(before) | set(after)
            if k != 'updated' and before.get(k) != after.get(k)}
    return r, diff, html


def find_browser():
    cands = [os.environ.get('TS_BROWSER', '')]
    for exe in ('google-chrome', 'google-chrome-stable', 'chromium', 'chromium-browser',
                'chrome', 'msedge', 'microsoft-edge'):
        cands.append(shutil.which(exe) or '')
    for base in (os.environ.get('PROGRAMFILES', ''), os.environ.get('PROGRAMFILES(X86)', ''),
                 os.environ.get('LOCALAPPDATA', '')):
        if base:
            cands += [os.path.join(base, 'Google', 'Chrome', 'Application', 'chrome.exe'),
                      os.path.join(base, 'Microsoft', 'Edge', 'Application', 'msedge.exe')]
    cands += ['/Applications/Google Chrome.app/Contents/MacOS/Google Chrome',
              '/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge']
    return next((c for c in cands if c and os.path.isfile(c)), None)


# The browser half's own routes, set up before the app serves anything.
from flask import request as _rq
RESULT = {}
HARNESS = r"""
var R = {};
function sleep(ms){ return new Promise(function(r){ setTimeout(r, ms); }); }
function frame(url){
  return new Promise(function(res){
    var f = document.createElement('iframe');
    f.style.cssText = 'width:1100px;height:800px;position:fixed;left:0;top:0';
    f.onload = function(){ f.onload = null; res(f); };
    f.src = url; document.body.appendChild(f);
  });
}
function nextLoad(f){ return new Promise(function(res){ f.onload = function(){ f.onload = null; res(f); }; }); }
function blocks(f){
  var w = f.contentWindow, ev = new w.Event('beforeunload', {cancelable: true});
  w.dispatchEvent(ev); return ev.defaultPrevented;
}
function retitle(f, t){
  var i = f.contentDocument.querySelector('input[name=title]');
  i.value = t; i.dispatchEvent(new f.contentWindow.Event('input', {bubbles: true}));
}
(async function(){
  try {
    // Customise with a changed title: saved, then on to the Wrap designer
    var f = await frame('/project/g1/edit');
    retitle(f, 'Changed Before Customise');
    var go = nextLoad(f);
    f.contentDocument.querySelector('a.custom[data-save-first]').click();
    await go;
    R.custom_path = f.contentWindow.location.pathname + f.contentWindow.location.search;
    // Write with nothing changed: just goes
    f = await frame('/project/g2/edit');
    go = nextLoad(f);
    f.contentDocument.querySelector('a[data-save-first][href*="/write"]').click();
    await go;
    R.write_path = f.contentWindow.location.pathname;
    // leaving with a change asks; unchanged doesn't; Cancel discards without asking
    f = await frame('/project/g3/edit');
    R.clean_blocks = blocks(f);
    retitle(f, 'Not Saved');
    R.dirty_blocks = blocks(f);
    go = nextLoad(f);
    f.contentDocument.getElementById('pe-cancel').click();
    R.cancel_blocks = blocks(f);
    await go;
  } catch(e){ R.error = String(e && e.stack || e); }
  await fetch('/_pe_result', {method: 'POST', body: JSON.stringify(R)});
})();
"""


@A.app.route('/_pe_harness')
def _pe_harness():
    return '<!doctype html><meta charset="utf-8"><body><script>' + HARNESS + '</script>'


@A.app.route('/_pe_result', methods=['POST'])
def _pe_result():
    RESULT.update(json.loads(_rq.get_data(as_text=True)))
    return 'ok'


def browser_half(tmp):
    print('\n[browser: leaving the page]')
    browser = find_browser()
    if not browser:
        print('SKIP  no Chrome/Edge found (set TS_BROWSER to a Chromium browser)')
        return
    from werkzeug.serving import make_server
    for pid in ('g1', 'g2', 'g3'):
        make_book(pid)
    before = {pid: stored(pid) for pid in ('g2', 'g3')}
    server = make_server('127.0.0.1', 0, A.app, threaded=True)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    try:
        # real time, as in test_manuscript_editor.py
        proc = subprocess.Popen([browser, '--headless=new', '--disable-gpu', '--no-first-run',
                                 '--remote-debugging-port=0',
                                 '--user-data-dir=' + os.path.join(tmp, 'browser'),
                                 f'http://127.0.0.1:{server.server_port}/_pe_harness'],
                                stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        try:
            for _ in range(240):
                if RESULT:
                    break
                time.sleep(0.5)
        finally:
            proc.kill()
            proc.wait()
    finally:
        server.shutdown()
    R = RESULT
    if not R:
        check('the browser ran the harness', False, browser)
        return
    check('the harness ran to the end', not R.get('error'), R.get('error'))
    check('Customise with unsaved changes saves them first',
          stored('g1')['title'] == 'Changed Before Customise')
    check('  and goes on to the Wrap designer', (R.get('custom_path') or '').startswith(
        '/wrap-designer') and 'from=' in R.get('custom_path', ''), R.get('custom_path'))
    check('Write with nothing changed goes straight there, saving nothing',
          (R.get('write_path') or '').endswith('/write') and stored('g2') == before['g2'],
          R.get('write_path'))
    check('leaving unchanged asks nothing', R.get('clean_blocks') is False)
    check('leaving with a change asks first', R.get('dirty_blocks') is True)
    check('Cancel leaves without asking, and saves nothing',
          R.get('cancel_blocks') is False and stored('g3') == before['g3'])


def main():
    tmp = tempfile.mkdtemp(prefix='ts-edit-')
    try:
        use_folders(tmp)
        c = A.app.test_client()
        cover = A.PROJECT_MS_DIR

        print('[an unchanged save is the same book]')
        with open(os.path.join(cover, 'img-cover.png'), 'wb') as f:
            f.write(png_bytes())
        with open(os.path.join(cover, 'img-back.png'), 'wb') as f:
            f.write(png_bytes(30, 30))
        cases = {
            'designed': {},
            'none': {'cover_mode': 'none'},
            'image': {'cover_mode': 'image', 'cover_file': 'img-cover.png',
                      'print_back_file': 'img-back.png'},
            'wrap': {'cover_mode': 'wrap', 'wrap_design': {'elements': []}},
            'awkward': {k: AWKWARD for k in A.matter.KEYS},
            'defaults': {k: v for k, v in A.PRINT_DEFAULTS.items()},
        }
        for name, over in cases.items():
            make_book(name, **over)
            r, diff, _ = unchanged_save(c, name)
            check(f'{name}: saved', r.status_code == 302, r.status_code)
            check(f'{name}: the same book', not diff, diff)

        print('\n[what the page has no option for]')
        shutil.copy(os.path.join(A.PRESET_DIR, 'classic-literary.json'),
                    os.path.join(A.PRESET_DIR, 'gone-style.json'))
        make_book('lost-style', preset='gone-style')
        os.remove(os.path.join(A.PRESET_DIR, 'gone-style.json'))
        r, diff, html = unchanged_save(c, 'lost-style')
        check('a deleted style is kept', 'preset' not in diff, diff.get('preset'))
        check('and the page says it is missing', 'gone-style' in html)

        make_book('lost-cover', cover_template='gone-cover')
        r, diff, html = unchanged_save(c, 'lost-cover')
        check('a deleted cover template is kept', 'cover_template' not in diff,
              diff.get('cover_template'))
        check('and the page says it is missing', 'gone-cover' in html)

        make_book('odd-print', print_retailer='lulu-2030', print_paper='vellum',
                  print_binding='spiral', format='mobi', front_matter='odd',
                  cover_color='plaid', cover_mode='hologram')
        r, diff, html = unchanged_save(c, 'odd-print')
        check('values this version has no option for are not swapped for the first',
              not diff, diff)

        print('\n[values]')
        make_book('nums')
        base = form_of(c.get('/project/nums/edit').get_data(as_text=True), '/project/nums/edit')
        for w, y, want_w, want_y in (('abc', '', 2.25, 0.7), ('nan', 'inf', 2.25, 0.7),
                                     ('-3', '7', 0.0, 1.0), ('1e308', '-1e308', None, 0.0)):
            f = base.copy()
            f['print_back_w'], f['print_back_y'] = w, y
            c.post('/project/nums/edit', data=f)
            got = stored('nums')
            ok_w = (got['print_back_w'] == want_w) if want_w is not None else \
                (0 <= got['print_back_w'] <= 20)
            check(f'back image width {w!r} / position {y!r} -> sensible',
                  ok_w and got['print_back_y'] == want_y,
                  (got['print_back_w'], got['print_back_y']))
            # always valid JSON (no NaN / Infinity in the file)
            with open(os.path.join(A.PROJECT_DIR, 'nums.json'), encoding='utf-8') as fh:
                txt = fh.read()
            check('  and the file is strict JSON', 'NaN' not in txt and 'Infinity' not in txt)

        f = base.copy(); f['name'] = '   '
        c.post('/project/nums/edit', data=f)
        check('a blank name keeps the old one', stored('nums')['name'] == 'Full Book')

        print('\n[uploads]')
        make_book('up')
        base = form_of(c.get('/project/up/edit').get_data(as_text=True), '/project/up/edit')

        def post(**files):
            f = base.copy()
            for k, v in files.items():
                f[k] = v
            return c.post('/project/up/edit', data=f, content_type='multipart/form-data',
                          follow_redirects=True)

        r = post(manuscript=(io.BytesIO(b'MZ\x90\x00binary'), 'virus.exe'))
        check('a manuscript that is not .docx/.md/.txt is refused',
              stored('up')['manuscript_file'] == 'up.md'
              and not os.path.exists(os.path.join(cover, 'up.exe')))
        check('  and the page says so', 'manuscript' in r.get_data(as_text=True).lower()
              and '.exe' in r.get_data(as_text=True))

        text = '# Kapitel\n\nGrüße aus Köln – “Anführung”.\n'
        r = post(manuscript=(io.BytesIO(text.encode('cp1252')), 'old.txt'))
        got = A._project_manuscript_text(stored('up'), report_import=False)
        check('a Windows-1252 .txt is read as written', got == text, repr(got[:60]))
        check('  the replaced draft is in History', any(
            A.read_snapshot('up', s['stamp']) == '# One\n\nThe text.\n'
            for s in A.list_snapshots('up')))

        r = post(manuscript=(io.BytesIO(b'\xef\xbb\xbf# Bom\n\nText.\n'), 'bom.md'))
        got = A._project_manuscript_text(stored('up'), report_import=False)
        parsed = manuscript.parse_markdown(got)
        check('a UTF-8 file with a byte-order mark keeps its first chapter',
              parsed['chapters'] and parsed['chapters'][0].get('title') == 'Bom',
              repr(got[:20]))

        r = post(manuscript=(io.BytesIO('# Wide\n\nText.\n'.encode('utf-16')), 'wide.txt'))
        got = A._project_manuscript_text(stored('up'), report_import=False)
        check('a UTF-16 .txt is read as written', got.startswith('# Wide'), repr(got[:20]))

        r = post(manuscript=(io.BytesIO(b'   \n\n'), 'blank.md'))
        check('an empty manuscript does not replace the text',
              A._project_manuscript_text(stored('up'), report_import=False).startswith('# Wide'))

        docx = io.BytesIO()
        import docx as _docx
        d = _docx.Document(); d.add_heading('Глава', 1); d.add_paragraph('Текст.'); d.save(docx)
        r = post(manuscript=(io.BytesIO(docx.getvalue()), 'Роман.docx'))
        got = A._project_manuscript_text(stored('up'), report_import=False)
        check('a Word file with a Cyrillic name is read as Word',
              stored('up')['manuscript_file'] == 'up.docx' and 'Текст.' in got, repr(got[:40]))
        r = post(manuscript=(io.BytesIO(b'PK\x03\x04 not really'), 'broken.docx'))
        check('a broken Word file is refused', 'Текст.' in A._project_manuscript_text(
            stored('up'), report_import=False) and 'Word' in r.get_data(as_text=True))

        r = post(cover=(io.BytesIO(png_bytes()), 'Обложка.png'))
        check('a cover with a Cyrillic name is taken', stored('up')['cover_file'] == 'up-cover.png',
              stored('up')['cover_file'])
        f = base.copy(); f['cover_mode'] = 'designed'
        os.remove(os.path.join(cover, 'up-cover.png'))
        A.save_project_file('up', dict(stored('up'), cover_file=''))

        r = post(cover=(io.BytesIO(b'GIF89a....'), 'cover.gif'))
        check('a cover that is not .jpg/.png is refused', stored('up')['cover_file'] == ''
              and 'cover' in r.get_data(as_text=True).lower())
        r = post(cover=(io.BytesIO(b'not an image at all'), 'cover.png'))
        check('a .png that is not an image is refused', stored('up')['cover_file'] == '',
              stored('up')['cover_file'])
        r = post(cover=(io.BytesIO(png_bytes()), 'cover.png'))
        check('a real cover is taken', stored('up')['cover_file'] == 'up-cover.png')

        r = post(print_back_image=(io.BytesIO(b'x'), 'me.bmp'))
        check('a back image of the wrong kind is refused, and said so',
              stored('up')['print_back_file'] == '' and '.bmp' in r.get_data(as_text=True))
        r = post(print_back_image=(io.BytesIO(png_bytes(20, 20)), 'me.png'))
        check('a back image is taken', stored('up')['print_back_file'] == 'up-back.png')
        f = base.copy(); f['print_back_clear'] = 'on'
        c.post('/project/up/edit', data=f)
        check('and cleared', stored('up')['print_back_file'] == ''
              and not os.path.exists(os.path.join(cover, 'up-back.png')))

        print('\n[the other uploads, with names in other scripts]')
        A.FIGURE_DIR = os.path.join(tmp, 'figures')
        A.FONT_DIR = os.path.join(tmp, 'fonts')
        os.makedirs(A.FIGURE_DIR); os.makedirs(A.FONT_DIR)
        names = []
        for n, colour in (('Карта.png', (10, 10, 10)), ('Схема.png', (200, 10, 10))):
            j = c.post('/figures/upload', data={'figures': (io.BytesIO(png_bytes(colour=colour)), n)},
                       content_type='multipart/form-data',
                       headers={'X-Requested-With': 'fetch'}).get_json()
            names += j['saved']
        check('two figures named in Cyrillic are both kept', len(set(names)) == 2
              and all(os.path.exists(os.path.join(A.FIGURE_DIR, n)) for n in names), names)
        with open(os.path.join(HERE, 'fonts', 'Alegreya-Regular.ttf'), 'rb') as fh:
            ttf = fh.read()
        c.post('/fonts/upload', data={'fonts': (io.BytesIO(ttf), 'Шрифт.ttf')},
               content_type='multipart/form-data')
        check('a font named in Cyrillic is taken',
              any(n.endswith('.ttf') for n in os.listdir(A.FONT_DIR)), os.listdir(A.FONT_DIR))

        browser_half(tmp)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    return 0 if not fails else 1


if __name__ == '__main__':
    code = main()
    print('\n' + ('ALL PASS' if not fails else 'FAILED: ' + ', '.join(fails)))
    sys.exit(code)
