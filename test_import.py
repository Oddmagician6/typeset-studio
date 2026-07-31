""".docx import fidelity tests  (run: python test_import.py).

The importer's failure mode is silence: `doc.paragraphs` skips whole element
types, and `Paragraph.runs` skips runs nested inside others, so things vanish
without anything going wrong. Every check here therefore builds a real Word
file containing the feature, imports it, and looks for the words on the other
side — and, for the ones that reach the page, builds the book too.

python-docx cannot author footnotes, so those parts are injected into the zip
by hand (`_with_notes`). That is the only way to exercise the real code path.
"""

import os
import sys
import json
import shutil
import tempfile
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass
import logging; logging.disable(logging.INFO)

import manuscript
import engine
import epub
import matter

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.enum.style import WD_STYLE_TYPE

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'

fails = []


def check(name, cond, detail=''):
    print(('  ok   ' if cond else '  FAIL ') + name +
          (('  ' + str(detail)) if not cond else ''))
    if not cond:
        fails.append(name)


# ---------------------------------------------------------------- fixture
def add_link(par, text, url):
    rid = par.part.relate_to(url, RT.HYPERLINK, is_external=True) if url else ''
    h = OxmlElement('w:hyperlink')
    if rid:
        h.set(qn('r:id'), rid)
    r = OxmlElement('w:r')
    t = OxmlElement('w:t')
    t.text = text
    r.append(t)
    h.append(r)
    par._p.append(h)


def add_note_ref(par, kind, nid):
    r = OxmlElement('w:r')
    ref = OxmlElement(f'w:{kind}Reference')
    ref.set(qn('w:id'), str(nid))
    r.append(ref)
    par._p.append(r)


def _note_part(tag, items):
    """A footnotes.xml / endnotes.xml part. Word reserves ids 0 and 1 for the
    separator rules it draws above the note area; the importer must skip them."""
    body = [f'<w:{tag} w:type="separator" w:id="0"><w:p/></w:{tag}>',
            f'<w:{tag} w:type="continuationSeparator" w:id="1"><w:p/></w:{tag}>']
    for nid, text in items:
        body.append(f'<w:{tag} w:id="{nid}"><w:p>'
                    f'<w:r><w:{tag}Ref/></w:r>'
                    f'<w:r><w:t xml:space="preserve"> {text}</w:t></w:r>'
                    f'</w:p></w:{tag}>')
    return ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<w:{tag}s xmlns:w="{W}">' + ''.join(body) + f'</w:{tag}s>')


FOOTNOTES = [(2, 'Ferrun Cartographic Office, Survey of the Salt Road, 1891.'),
             (3, 'The disagreement was never formally settled.'),
             (4, 'A note belonging to the second chapter.')]
ENDNOTES = [(2, 'An endnote, which Word keeps in a different part entirely.')]

_CT = ('<Override PartName="/word/footnotes.xml" ContentType="application/vnd.'
       'openxmlformats-officedocument.wordprocessingml.footnotes+xml"/>'
       '<Override PartName="/word/endnotes.xml" ContentType="application/vnd.'
       'openxmlformats-officedocument.wordprocessingml.endnotes+xml"/>')
_RELS = ('<Relationship Id="rIdFootnotes" Type="http://schemas.openxmlformats.org/'
         'officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/>'
         '<Relationship Id="rIdEndnotes" Type="http://schemas.openxmlformats.org/'
         'officeDocument/2006/relationships/endnotes" Target="endnotes.xml"/>')


def _with_notes(src, dest):
    with zipfile.ZipFile(src) as zin, \
            zipfile.ZipFile(dest, 'w', zipfile.ZIP_DEFLATED) as zo:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == '[Content_Types].xml':
                data = data.replace(b'</Types>', _CT.encode() + b'</Types>')
            elif item.filename == 'word/_rels/document.xml.rels':
                data = data.replace(b'</Relationships>',
                                    _RELS.encode() + b'</Relationships>')
            zo.writestr(item, data)
        zo.writestr('word/footnotes.xml', _note_part('footnote', FOOTNOTES))
        zo.writestr('word/endnotes.xml', _note_part('endnote', ENDNOTES))


def build_fixture(path):
    d = Document()
    d.add_heading('The Salt Road', level=1)

    p = d.add_paragraph('The plateau was surveyed twice.')
    add_note_ref(p, 'footnote', 2)
    p.add_run(' The rolls are held by ')
    add_link(p, 'the survey office', 'https://example.com/survey_a')
    p.add_run('. Write to ')
    add_link(p, 'the clerk', 'mailto:clerk@example.com')
    p.add_run(', or ask about ')
    add_link(p, 'the bookmark', '')            # no address: words only
    p.add_run('.')

    p = d.add_paragraph('The second survey disagreed.')
    add_note_ref(p, 'endnote', 2)
    p.add_run(' It was never settled.')
    add_note_ref(p, 'footnote', 3)

    p = d.add_paragraph('Ferrun Cartographic Office')     # manual line breaks
    p.add_run().add_break(); p.add_run('12 Rope Street')
    p.add_run().add_break(); p.add_run('Northmarch')

    d.styles.add_style('Verse', WD_STYLE_TYPE.PARAGRAPH)
    v = d.add_paragraph(style='Verse')
    v.add_run('Two roads diverged in a yellow wood,')
    v.add_run().add_break(); v.add_run('And sorry I could not travel both')
    v = d.add_paragraph(style='Verse')
    v.add_run('Then took the other, as just as fair,')
    v.add_run().add_break(); v.add_run('And ')
    v.add_run('having').italic = True
    v.add_run(' perhaps the better claim,')

    t = d.add_table(rows=2, cols=2)
    for r, row in enumerate([['Region', 'Wheat'], ['Northmarch', '1,240']]):
        for c, val in enumerate(row):
            t.rows[r].cells[c].text = val

    d.add_heading('Second Piece', level=1)
    p = d.add_paragraph('A later chapter cites one too.')
    add_note_ref(p, 'footnote', 4)

    base = path + '.base'
    d.save(base)
    _with_notes(base, path)
    os.remove(base)


META = {'title': 'Import', 'subtitle': '', 'author': 'E. Vale', 'year': '2026',
        'publisher': '', 'front_matter': 'none', 'right_hand_starts': False,
        'include_toc': False, 'smartquotes': True, 'cover_mode': 'none',
        'cover_image': '', 'cover_overlay': False, 'cover_color': 'light',
        **matter.blank()}


# ---------------------------------------------------------------- the run
tmp = tempfile.mkdtemp(prefix='tsimport-')
figdir = os.path.join(tmp, 'figures')
os.makedirs(figdir, exist_ok=True)
_saved_figdir = manuscript.FIGURE_DIR
manuscript.FIGURE_DIR = figdir
docx_path = os.path.join(tmp, 'salt-road.docx')

try:
    build_fixture(docx_path)
    report = {}
    md = manuscript.import_docx(docx_path, report=report)
    parsed = manuscript.parse_markdown(md, smartquotes=True)
    ch1, ch2 = parsed['chapters']

    print('notes')
    # a Word footnote and a Word endnote are different parts with overlapping
    # ids, so keying on the id alone would collapse two notes into one
    check('every note is imported', report['notes'] == 4, report['notes'])
    check('none was unreadable', report['notes_lost'] == 0, report['notes_lost'])
    check('the reference sits in the sentence, not at the end',
          'twice.<note' in ch1['blocks'][0][1], ch1['blocks'][0][1][:80])
    check('chapter one holds its three notes', len(ch1['notes']) == 3,
          len(ch1.get('notes', [])))
    check('the footnote text came across',
          ch1['notes'][0]['text'].startswith('Ferrun Cartographic Office'),
          ch1['notes'][0]['text'])
    check('the endnote part was read too, not just footnotes.xml',
          'Word keeps in a different part' in ch1['notes'][1]['text'],
          ch1['notes'][1]['text'])
    check("Word's separator pseudo-notes are skipped",
          all('separator' not in n['text'] for n in ch1['notes']))
    check('a note lands in the chapter that cites it',
          len(ch2['notes']) == 1 and 'second chapter' in ch2['notes'][0]['text'],
          ch2.get('notes'))
    check('numbering restarts per chapter',
          [n['n'] for n in ch1['notes']] == [1, 2, 3]
          and [n['n'] for n in ch2['notes']] == [1],
          ([n['n'] for n in ch1['notes']], [n['n'] for n in ch2['notes']]))

    print('links')
    body = ch1['blocks'][0][1]
    check('a web address survives',
          '<a href="https://example.com/survey_a">the survey office</a>' in body, body)
    check('a mailto address survives',
          '<a href="mailto:clerk@example.com">the clerk</a>' in body, body)
    check('an underscore in the URL is not read as emphasis',
          'survey_a' in body and '<i>' not in body, body)
    check('a link with no usable address keeps its words',
          'ask about the bookmark' in body, body)
    check('and is reported as such', report['links'] == 1 and report['links_kept'] == 2,
          (report['links'], report['links_kept']))

    print('verse and line breaks')
    kinds = [b[2].get('_type') for b in ch1['blocks'] if b[0] == 'doc_block']
    check('the block order is address, poem, table',
          kinds == ['left', 'poem', 'table'], kinds)
    poem = [b for b in ch1['blocks']
            if b[0] == 'doc_block' and b[2].get('_type') == 'poem'][0]
    check('two Word paragraphs become two stanzas', len(poem[1]) == 2, len(poem[1]))
    check('verse lines are preserved, not joined',
          poem[1][0][1].count('<br/>') == 1, poem[1][0][1])
    check('emphasis inside verse survives', '<i>having</i>' in poem[1][1][1],
          poem[1][1][1])
    addr = [b for b in ch1['blocks']
            if b[0] == 'doc_block' and b[2].get('_type') == 'left'][0]
    check('manual line breaks are kept as three lines', len(addr[1]) == 3,
          [t for _, t in addr[1]])
    check('and are reported', report['breaks'] == 1, report['breaks'])

    print('regressions from earlier import work')
    table = [b for b in ch1['blocks']
             if b[0] == 'doc_block' and b[2].get('_type') == 'table'][0]
    check('a Word table still imports with its columns (#57)',
          table[1][0][1].split(manuscript.CELL_SEP) == ['Region', 'Wheat'],
          table[1][0][1])
    check('both chapters are found', (ch1['title'], ch2['title'])
          == ('The Salt Road', 'Second Piece'), (ch1['title'], ch2['title']))

    print('the summary tells the truth')
    got, lost = manuscript.import_summary(report)
    for phrase in ('2 chapters', '1 poem', '2 links', '4 notes (as endnotes)',
                   '1 table', 'kept line breaks'):
        check(f'summary mentions {phrase!r}', phrase in got, got)
    check('the only loss reported is the address-less link',
          lost == '1 link kept its text but not the address', lost)

    print('buffered blocks keep document order')
    # a run of list paragraphs is buffered until something ends it. The table
    # branch used to flush only the quote buffer, so a list (or a poem) sitting
    # right before a table came out *after* it.
    d2 = Document()
    d2.add_heading('Order', level=1)
    d2.add_paragraph('Salt, four measures', style='List Bullet')
    d2.add_paragraph('Rope, two coils', style='List Bullet')
    t2 = d2.add_table(rows=1, cols=2)
    t2.rows[0].cells[0].text = 'A'
    t2.rows[0].cells[1].text = 'B'
    d2.add_paragraph('After the table.')
    order_path = os.path.join(tmp, 'order.docx')
    d2.save(order_path)
    ordered = manuscript.parse_markdown(manuscript.import_docx(order_path))
    seq = [b[2].get('_type') if b[0] == 'doc_block' else b[0]
           for b in ordered['chapters'][0]['blocks']]
    check('a list before a table stays before it',
          seq == ['list', 'table', 'para'], seq)

    print('re-import is idempotent')
    md2 = manuscript.import_docx(docx_path)
    check('a second import is byte-identical', md == md2)

    print('the imported book builds')
    preset = json.load(open(os.path.join(HERE, 'presets', 'science-fiction-clean.json'),
                            encoding='utf-8'))     # open_style "none": see below
    pdf = os.path.join(tmp, 'out.pdf')
    result = engine.build_pdf(parsed, preset, pdf, dict(META))
    check('the PDF builds with no note left unplaced',
          result.get('notes_unplaced', 0) == 0, result.get('notes_unplaced'))
    try:
        import fitz
        doc = fitz.open(pdf)
        uris = sorted({l['uri'] for pg in doc for l in pg.get_links() if l.get('uri')})
        text = '\n'.join(pg.get_text() for pg in doc)
        doc.close()
        # a link in a chapter's *first* paragraph is lost under the decorative
        # opening styles (a documented limit since #49), so this preset opens plain
        check('both addresses are live links in the PDF',
              uris == ['https://example.com/survey_a', 'mailto:clerk@example.com'], uris)
        check('the Notes page carries the note text',
              'Survey of the Salt Road' in text)
        check('the verse reached the page', 'yellow wood' in text)
    except ImportError:
        print('  --   pymupdf absent, skipping the PDF inspection')

    ep = os.path.join(tmp, 'out.epub')
    epub.build_epub(parsed, preset, ep, dict(META))
    bad = [c['label'] for c in epub.check(ep) if not c['ok']]
    check('the EPUB passes preflight', not bad, bad)
    with zipfile.ZipFile(ep) as zf:
        xhtml = zf.read('OEBPS/chapter001.xhtml').decode()
    check('the EPUB carries the real links',
          'href="mailto:clerk@example.com"' in xhtml, xhtml[:0])
finally:
    manuscript.FIGURE_DIR = _saved_figdir
    shutil.rmtree(tmp, ignore_errors=True)

# ---------------------------------------------------------------------------
# Word prose is prose, not Markdown.
#
# Everything above checks that a feature survives the trip. These check the
# opposite failure: text that was never markup arriving as markup. A Word
# paragraph beginning "#" opened a new chapter, one beginning "~~~" swallowed
# the rest of the document into a block, and a writer's literal *stars* and
# file_name came back italicised.
# ---------------------------------------------------------------------------
print('\n[word prose is not markdown]')
_tmp2 = tempfile.mkdtemp()
try:
    def _imported(build):
        path = os.path.join(_tmp2, 'p.docx')
        d = Document()
        build(d)
        d.save(path)
        rep = {}
        md = manuscript.import_docx(path, report=rep)
        return md, rep, manuscript.parse_markdown(md)

    def _blocks(parsed):
        return [b for c in parsed['chapters'] for b in c['blocks']]

    def _markers(d):
        d.add_heading('Real Chapter', level=1)
        d.add_paragraph('Ordinary opening line.')
        d.add_paragraph('#1 bestseller, said the blurb')
        d.add_paragraph('~~~ and this looked like a fence')
        d.add_paragraph('=== and this like a part divider')
        d.add_paragraph('She said *hello* with literal asterisks')
        d.add_paragraph('The file_name_here has underscores')

    md, rep, parsed = _imported(_markers)
    check('a line starting "#" does not open a chapter',
          len(parsed['chapters']) == 1, [c['title'] for c in parsed['chapters']])
    texts = [b[1] for b in _blocks(parsed) if b[0] == 'para']
    check('the "#" line stays body text',
          any(t.startswith('#1 bestseller') for t in texts), texts)
    check('a "~~~" line does not open a block',
          all(b[0] != 'doc_block' for b in _blocks(parsed)), _blocks(parsed))
    check('a "===" line stays body text',
          any(t.startswith('=== and this') for t in texts), texts)
    check('literal asterisks are not italicised',
          any('*hello*' in t and '<i>' not in t for t in texts), texts)
    check('an underscored word is not italicised',
          any('file_name_here' in t for t in texts), texts)

    def _bold_heading(d):
        h = d.add_heading('', level=1)
        r = h.add_run('A Bold Title')
        r.bold = True
        d.add_paragraph('Body.')

    md, rep, parsed = _imported(_bold_heading)
    check('a bolded Word heading gives a clean title',
          parsed['chapters'][0]['title'] == 'A Bold Title',
          repr(parsed['chapters'][0]['title']))

    def _bold_italic(d):
        d.add_heading('C', level=1)
        d.add_paragraph('Opening line.')
        p = d.add_paragraph()
        r = p.add_run('utterly certain')
        r.bold = True
        r.italic = True

    md, rep, parsed = _imported(_bold_italic)
    check('a bold+italic Word run keeps both',
          any('<b><i>utterly certain</i></b>' in b[1]
              for b in _blocks(parsed) if b[0] == 'para'),
          [b[1] for b in _blocks(parsed)])

    def _scene(d):
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        d.add_heading('C', level=1)
        d.add_paragraph('Before.')
        sb = d.add_paragraph('* * *')
        sb.alignment = WD_ALIGN_PARAGRAPH.CENTER
        d.add_paragraph('After.')

    md, rep, parsed = _imported(_scene)
    check('a scene break still imports as a scene break',
          any(b[0] == 'scene' for b in _blocks(parsed)), _blocks(parsed))

    def _nested(d):
        d.add_heading('C', level=1)
        d.add_paragraph('Body.')
        t = d.add_table(rows=1, cols=2)
        t.cell(0, 0).text = 'outer'
        inner = t.cell(0, 1).add_table(rows=1, cols=2)
        inner.cell(0, 0).text = 'INNER-ONE'
        inner.cell(0, 1).text = 'INNER-TWO'

    md, rep, parsed = _imported(_nested)
    check('a nested table is not dropped', rep.get('tables') == 1, rep)
    check('and its words reach the book',
          'INNER-ONE' in md and 'INNER-TWO' in md, md)

    def _markup_cells(d):
        d.add_heading('C', level=1)
        d.add_paragraph('Body.')
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text = 'Name'
        t.cell(0, 1).text = 'Note'
        t.cell(1, 0).text = 'file_name'
        t.cell(1, 1).text = '*starred*'

    md, rep, parsed = _imported(_markup_cells)
    cell_text = ' '.join(line[1]
                         for block in _blocks(parsed) if block[0] == 'doc_block'
                         for line in block[1])
    check('the table block carries the cells', 'Name' in cell_text, cell_text)
    check('markup characters in a cell stay literal',
          'file_name' in cell_text and '*starred*' in cell_text
          and '<i>' not in cell_text, cell_text)
finally:
    shutil.rmtree(_tmp2, ignore_errors=True)

print('\n' + ('ALL PASS' if not fails else f'FAILED ({len(fails)}): ' + ', '.join(fails)))
sys.exit(1 if fails else 0)
