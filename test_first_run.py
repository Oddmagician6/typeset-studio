"""Setting a book, the first-run path  (run: python test_first_run.py).

Bug hunt item 4: a new user's first ten minutes. Set a book (/generate), its
two previews, saving the result as a project (/project/create), a blank new
draft, an empty data folder. Checked:

    ROUND TRIP  a book Set with every option, then saved as a project, builds as
                the same book (every page's text), from an upload, pasted text,
                the sample, and with cover art
    PREVIEWS    both previews take every source, and the page preview shows the
                pages Set a book builds, pixel for pixel
    UPLOADS     wrong kinds, empty, other encodings, names in other scripts, the
                same file name twice before saving, a huge manuscript
    SAVING      only uploads become a project's files; ids and file names for
                titles in other scripts
    FIRST RUN   no styles at all, no books at all, a blank draft
    BROWSER     the real page: choose a file, both previews, Set, Save as project
"""

import sys, os, io, json, time, shutil, tempfile, threading, subprocess
from html.parser import HTMLParser

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass
import logging; logging.disable(logging.INFO)

import app as A

fails = []
def check(name, cond, detail=''):
    print(('  ok   ' if cond else '  FAIL ') + name + (('  ' + str(detail)) if not cond else ''))
    if not cond:
        fails.append(name)


class FormFields(HTMLParser):
    """What a browser would submit from the form posting to `action`, untouched.
    (As in test_upgrade.py.)"""
    def __init__(self, action):
        super().__init__(convert_charrefs=True)
        self.action, self.inside, self.fields = action, False, []
        self.select, self.options, self.textarea = None, [], None

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if tag == 'form':
            self.inside = self.action in (a.get('action') or '')
        if not self.inside or 'disabled' in a:
            return
        if tag == 'input' and a.get('name'):
            kind = (a.get('type') or 'text').lower()
            if kind in ('checkbox', 'radio'):
                if 'checked' in a:
                    self.fields.append((a['name'], a.get('value', 'on')))
            elif kind not in ('file', 'submit', 'button', 'image', 'reset'):
                self.fields.append((a['name'], a.get('value', '')))
        elif tag == 'select' and a.get('name'):
            self.select, self.options = a['name'], []
        elif tag == 'option' and self.select:
            self.options.append([a.get('value'), 'selected' in a, ''])
        elif tag == 'textarea' and a.get('name'):
            self.textarea = [a['name'], '']

    def handle_data(self, data):
        if self.textarea is not None:
            self.textarea[1] += data
        elif self.select and self.options and self.options[-1][0] is None:
            self.options[-1][2] += data

    def handle_endtag(self, tag):
        if tag == 'form':
            self.inside = False
        elif tag == 'textarea' and self.textarea is not None:
            name, text = self.textarea
            self.fields.append((name, text[1:] if text.startswith('\n') else text))
            self.textarea = None
        elif tag == 'select' and self.select:
            opts = [(v if v is not None else t.strip(), s) for v, s, t in self.options]
            pick = next((v for v, s in opts if s), opts[0][0] if opts else None)
            if pick is not None:
                self.fields.append((self.select, pick))
            self.select = None


def form_of(html, action):
    from werkzeug.datastructures import MultiDict
    p = FormFields(action)
    p.feed(html)
    return MultiDict(p.fields)


def find_browser():
    cands = [os.environ.get('TS_BROWSER', '')]
    for exe in ('google-chrome', 'google-chrome-stable', 'chromium', 'chromium-browser',
                'chrome', 'msedge', 'microsoft-edge'):
        cands.append(shutil.which(exe) or '')
    for base in (os.environ.get('PROGRAMFILES', ''), os.environ.get('PROGRAMFILES(X86)', ''),
                 os.environ.get('LOCALAPPDATA', '')):
        if base:
            cands += [os.path.join(base, 'Google', 'Chrome', 'Application', 'chrome.exe'),
                      os.path.join(base, 'Microsoft', 'Edge', 'Application', 'msedge.exe')]
    cands += ['/Applications/Google Chrome.app/Contents/MacOS/Google Chrome',
              '/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge']
    return next((c for c in cands if c and os.path.isfile(c)), None)


def use_folders(root, styles=True):
    os.makedirs(root)
    if styles:
        shutil.copytree(os.path.join(HERE, 'presets'), os.path.join(root, 'presets'))
    else:
        os.makedirs(os.path.join(root, 'presets'))
    shutil.copytree(os.path.join(HERE, 'covers'), os.path.join(root, 'covers'),
                    ignore=shutil.ignore_patterns('assets'))
    A.PRESET_DIR, A.COVER_DIR = os.path.join(root, 'presets'), os.path.join(root, 'covers')
    A.PROJECT_DIR = os.path.join(root, 'projects')
    A.PROJECT_MS_DIR = os.path.join(A.PROJECT_DIR, 'manuscripts')
    A.HISTORY_DIR = os.path.join(A.PROJECT_DIR, 'history')
    A.OUT_DIR = os.path.join(root, 'out')
    A.UPLOAD_DIR = os.path.join(root, 'uploads')
    A.COVER_THUMB_DIR = os.path.join(A.OUT_DIR, '_cover_thumbs')
    A.PROJECT_THUMB_DIR = os.path.join(A.OUT_DIR, '_project_thumbs')
    for d in (A.PROJECT_MS_DIR, A.HISTORY_DIR, A.UPLOAD_DIR, A.COVER_THUMB_DIR,
              A.PROJECT_THUMB_DIR):
        os.makedirs(d, exist_ok=True)


def png_bytes(w=600, h=900, colour=(40, 60, 90)):
    from PIL import Image
    b = io.BytesIO()
    Image.new('RGB', (w, h), colour).save(b, 'PNG')
    return b.getvalue()


def docx_bytes(paras):
    import docx
    d = docx.Document()
    d.add_heading(paras[0], 1)
    for p in paras[1:]:
        d.add_paragraph(p)
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def pdf_pages(path, zoom=None):
    import fitz
    doc = fitz.open(path)
    if zoom:
        out = [p.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False).tobytes('png')
               for p in doc]
    else:
        out = [p.get_text() for p in doc]
    doc.close()
    return out


with open(os.path.join(HERE, 'sample', 'sample.md'), encoding='utf-8') as _f:
    SAMPLE = _f.read()

# Every option the compose form offers, set to something other than its default.
FULL = {
    'preset': 'classic-literary', 'format': 'both',
    'title': 'The Salt Road', 'subtitle': 'A Novel', 'author': 'Ellinor Vale',
    'year': '2031', 'publisher': 'Ashforge & Sons', 'front_matter': 'title',
    'right_hand_starts': 'on', 'include_toc': 'on', 'smartquotes': 'on',
    'cover_mode': 'designed', 'cover_template': 'literary-ivory',
    'cover_collection': 'The Salt Cycle', 'cover_kicker': 'A tale of the coast',
    'cover_accent': 'E. Vale', 'cover_epigraph': 'The sea keeps what it takes.',
    'cover_studio': 'Ashforge', 'cover_color': 'dark',
}


def compose(c, fields=None, files=None, follow=False, url='/generate'):
    data = dict(fields if fields is not None else FULL)
    for k, v in (files or {}).items():
        data[k] = (io.BytesIO(v[1]), v[0])
    return c.post(url, data=data, content_type='multipart/form-data',
                  follow_redirects=follow)


def save_result(c, html, name=None):
    form = form_of(html, '/project/create')
    if name is not None:
        form['project_name'] = name
    before = set(os.listdir(A.PROJECT_DIR))
    c.post('/project/create', data=form)
    new = sorted(set(os.listdir(A.PROJECT_DIR)) - before)
    new = [n for n in new if n.endswith('.json')]
    return (new[0][:-5] if new else None), form


def load(pid):
    with open(os.path.join(A.PROJECT_DIR, pid + '.json'), encoding='utf-8') as f:
        return json.load(f)


def round_trip(c, label, fields, files=None):
    """Set a book, save it as a project, build the project: the same pages?"""
    r = compose(c, fields, files)
    html = r.get_data(as_text=True)
    if r.status_code != 200 or '/project/create' not in html:
        check(f'{label}: Set a book builds', False, (r.status_code, html[:300]))
        return None
    pid, form = save_result(c, html)
    if not pid:
        check(f'{label}: saved as a project', False)
        return None
    proj = load(pid)
    res = A._build_project(pid, proj)
    if res.get('error'):
        check(f'{label}: the project builds', False, res['error'])
        return pid
    proj = load(pid)
    set_pdf = os.path.join(A.OUT_DIR, form['last_pdf'])
    built = os.path.join(A.OUT_DIR, proj['last_pdf'])
    a, b = pdf_pages(set_pdf), pdf_pages(built)
    year_fix = [t.replace(str(time.localtime().tm_year), 'YEAR') for t in a]
    check(f'{label}: the project builds the same book ({len(a)} pages)',
          a == b or year_fix == [t.replace(str(time.localtime().tm_year), 'YEAR') for t in b],
          next(((i, x[:80], y[:80]) for i, (x, y) in enumerate(zip(a, b)) if x != y),
               (len(a), len(b))))
    return pid


# The browser half's routes, registered before the app serves a request.
from flask import request as _rq
RESULT = {}
HARNESS = r"""
var R = {};
function sleep(ms){ return new Promise(function(r){ setTimeout(r, ms); }); }
function frame(url){
  return new Promise(function(res){
    var f = document.createElement('iframe');
    f.style.cssText = 'width:1200px;height:900px;position:fixed;left:0;top:0';
    f.onload = function(){ f.onload = null; res(f); };
    f.src = url; document.body.appendChild(f);
  });
}
function nextLoad(f){ return new Promise(function(res){ f.onload = function(){ f.onload = null; res(f); }; }); }
async function until(test, ms){
  for (var t = 0; t < ms; t += 100){ if (test()) return true; await sleep(100); }
  return false;
}
(async function(){
  try {
    var f = await frame('/generate'), d = f.contentDocument, w = f.contentWindow;
    var dt = new w.DataTransfer();
    dt.items.add(new w.File(['# Глава первая\n\nТекст первой главы.\n\n# Second\n\nMore text.\n'],
                            'Роман.md', {type: 'text/markdown'}));
    d.querySelector('input[name=manuscript]').files = dt.files;
    d.querySelector('input[name=title]').value = 'Война';
    d.getElementById('pv-btn').click();
    R.pv_shown = await until(function(){ return d.querySelectorAll('#pv-body img').length > 0; }, 20000);
    R.pv_cap = d.getElementById('pv-cap').textContent;
    d.getElementById('pv-close').click();
    d.getElementById('ev-btn').click();
    R.ev_shown = await until(function(){ return !!d.querySelector('#ev-body iframe'); }, 20000);
    var ef = d.querySelector('#ev-body iframe');
    R.ev_has_text = !!ef && ef.srcdoc.indexOf('Текст первой главы') >= 0;
    d.getElementById('ev-close').click();
    var go = nextLoad(f);
    d.querySelector('form.compose').requestSubmit();
    await go;
    d = f.contentDocument;
    R.result_has_save = !!d.querySelector('form[action$="/project/create"]');
    go = nextLoad(f);
    d.querySelector('form[action$="/project/create"]').requestSubmit();
    await go;
    R.after_save = f.contentWindow.location.pathname;
  } catch(e){ R.error = String(e && e.stack || e); }
  await fetch('/_fr_result', {method: 'POST', body: JSON.stringify(R)});
})();
"""


@A.app.route('/_fr_harness')
def _fr_harness():
    return '<!doctype html><meta charset="utf-8"><body><script>' + HARNESS + '</script>'


@A.app.route('/_fr_result', methods=['POST'])
def _fr_result():
    RESULT.update(json.loads(_rq.get_data(as_text=True)))
    return 'ok'


def main():
    tmp = tempfile.mkdtemp(prefix='ts-first-')
    try:
        use_folders(os.path.join(tmp, 'data'))
        c = A.app.test_client()

        print('[round trip: Set a book, save it, build the project]')
        round_trip(c, 'uploaded .md', FULL, {'manuscript': ('salt-road.md', SAMPLE.encode())})
        round_trip(c, 'pasted text', dict(FULL, pasted=SAMPLE, format='pdf'))
        round_trip(c, 'the sample', dict(FULL, use_sample='on', format='pdf',
                                         cover_mode='none'))
        pid = round_trip(c, 'cover art', dict(FULL, cover_mode='image', cover_overlay='on',
                                              format='pdf'),
                         {'manuscript': ('art.md', SAMPLE.encode()),
                          'cover': ('front.png', png_bytes())})
        if pid:
            check('  the cover art is the project\'s own file', load(pid)['cover_file'] == pid + '-cover.png'
                  and os.path.exists(os.path.join(A.PROJECT_MS_DIR, pid + '-cover.png')))
        pid = round_trip(c, 'a Word file named in Cyrillic',
                         dict(FULL, format='pdf', title='Война и мир'),
                         {'manuscript': ('Роман.docx', docx_bytes(['Глава', 'Текст главы.']))})
        if pid:
            check('  its project id and files are named sensibly',
                  pid.startswith('book') and load(pid)['last_pdf'].startswith('book-'),
                  (pid, load(pid)['last_pdf']))
        r = compose(c, dict(FULL, format='pdf', title='Été à Paris', cover_mode='none'),
                    {'manuscript': ('a.md', b'# One\n\nText.\n')})
        check('accents in a title are folded, not dropped',
              'ete-a-paris-' in r.get_data(as_text=True))

        print('\n[uploads]')
        for name, data, want in (
                ('book.pdf', b'%PDF-1.4 nonsense', "isn't a manuscript"),
                ('empty.md', b'  \n\n', 'is empty'),
                ('broken.docx', b'PK\x03\x04 nope', "couldn't be read as a Word"),
                ('notes.docx', docx_bytes(['']), 'has no text')):
            r = compose(c, dict(FULL, format='pdf'), {'manuscript': (name, data)})
            html = r.get_data(as_text=True).replace('&#39;', "'").replace('&#x27;', "'")
            check(f'{name}: refused, and says why', want in html and '/project/create' not in html)
        r = compose(c, dict(FULL, format='pdf', cover_mode='image'),
                    {'manuscript': ('a.md', b'# One\n\nText.\n'), 'cover': ('c.png', b'no')})
        check('cover art that is not a picture: refused', "picture" in r.get_data(as_text=True)
              and '/project/create' not in r.get_data(as_text=True))
        r = compose(c, dict(FULL, format='pdf', cover_mode='designed'),
                    {'manuscript': ('a.md', b'# One\n\nText.\n'), 'cover': ('c.png', b'no')})
        check('  but ignored when the book uses a designed cover',
              '/project/create' in r.get_data(as_text=True))
        text = '# Kapitel\n\nGrüße aus Köln – “Anführung”.\n'
        r = compose(c, dict(FULL, format='pdf', smartquotes=''),
                    {'manuscript': ('old.txt', text.encode('cp1252'))})
        pdf = form_of(r.get_data(as_text=True), '/project/create')['last_pdf']
        got = ''.join(pdf_pages(os.path.join(A.OUT_DIR, pdf)))
        # the opener sets the first words in small capitals
        check('a Windows-1252 .txt is set as written',
              'GRÜSSE AUS KÖLN' in got and '“Anführung”' in got, got[-80:])

        # the same file name twice, the first result saved after the second upload
        r1 = compose(c, dict(FULL, format='pdf'), {'manuscript': ('book.md', b'# First\n\nFIRST BOOK.\n')})
        r2 = compose(c, dict(FULL, format='pdf'), {'manuscript': ('book.md', b'# Second\n\nSECOND BOOK.\n')})
        pid, _ = save_result(c, r1.get_data(as_text=True), name='first of two')
        got = A._project_manuscript_text(load(pid), report_import=False)
        check('two uploads with one name: each result keeps its own', 'FIRST BOOK' in got, got[:40])

        huge = ''.join(f'# Chapter {i}\n\n' + ('A long paragraph of the book. ' * 60 + '\n\n') * 30
                       for i in range(1, 80))
        t0 = time.time()
        j = compose(c, dict(FULL, format='pdf'), {'manuscript': ('huge.md', huge.encode())},
                    url='/generate/preview').get_json()
        check(f'a {len(huge) // 1024} KB manuscript previews quickly ({time.time() - t0:.1f}s)',
              j['ok'] and j['truncated'] and time.time() - t0 < 20, j.get('error'))

        print('\n[saving as a project]')
        r = compose(c, dict(FULL, format='pdf'), {'manuscript': ('x.md', b'# X\n\nText.\n')})
        form = form_of(r.get_data(as_text=True), '/project/create')
        secret = os.path.join(tmp, 'private.txt')
        with open(secret, 'w') as f:
            f.write('# Not yours\n\nPRIVATE.\n')
        form['ms_path'], form['cover_path'], form['project_name'] = secret, secret, 'Sneaky'
        c.post('/project/create', data=form)
        proj = load('sneaky')
        check('a path outside the uploads folder is not copied into a book',
              proj['manuscript_file'] == '' and proj['cover_file'] == '', proj)

        print('\n[previews]')
        for label, fields, files in (
                ('upload', dict(FULL), {'manuscript': ('p.md', SAMPLE.encode())}),
                ('paste', dict(FULL, pasted=SAMPLE), None),
                ('sample', dict(FULL, use_sample='on'), None),
                ('Word, Cyrillic name', dict(FULL), {'manuscript': ('Книга.docx', docx_bytes(['Глава', 'Текст.']))})):
            j = compose(c, fields, files, url='/generate/preview').get_json()
            check(f'page preview from {label}', j['ok'] and j['images'], j.get('error'))
            j = compose(c, fields, files, url='/generate/epub-preview').get_json()
            check(f'ebook preview from {label}', j['ok'] and j['docs'], j.get('error'))
        j = compose(c, {k: v for k, v in FULL.items() if k != 'preset'},
                    {'manuscript': ('p.md', SAMPLE.encode())}, url='/generate/preview').get_json()
        check('no style: the preview says to pick one', not j['ok'] and 'style' in j['error'])
        j = compose(c, dict(FULL), url='/generate/preview').get_json()
        check('no manuscript: the preview says to add one', not j['ok'] and 'manuscript' in j['error'])

        # the page preview is the pages Set a book builds
        short = '# One\n\nThe first chapter.\n\n# Two\n\nThe second.\n'
        fields = dict(FULL, format='pdf', pasted=short)
        j = compose(c, fields, url='/generate/preview').get_json()
        r = compose(c, fields)
        pdf = form_of(r.get_data(as_text=True), '/project/create')['last_pdf']
        import base64
        built = pdf_pages(os.path.join(A.OUT_DIR, pdf), zoom=2)
        shown = [base64.b64decode(u.split(',', 1)[1]) for u in j['images']]
        check(f'the page preview is the book Set a book builds ({len(shown)} pages)',
              shown == built[:len(shown)] and len(shown) == min(len(built), A.PREVIEW_MAX_PAGES))

        print('\n[first run]')
        use_folders(os.path.join(tmp, 'empty'), styles=False)
        for url in ('/', '/generate', '/projects', '/covers'):
            r = c.get(url)
            check(f'no styles, no books: {url} opens', r.status_code == 200, r.status_code)
        r = c.post('/project/new-draft', data={'name': ''})
        pid = r.headers['Location'].rstrip('/').split('/')[-2]
        check('a blank draft opens the writing page', c.get(r.headers['Location']).status_code == 200)
        r = c.post(f'/project/{pid}/generate', follow_redirects=True)
        check('  and typesetting it with no style says what to do, not a crash',
              r.status_code == 200 and 'style' in r.get_data(as_text=True).lower(), r.status_code)
        use_folders(os.path.join(tmp, 'fresh'))
        r = c.post('/project/new-draft', data={'name': 'Роман'})
        pid = r.headers['Location'].rstrip('/').split('/')[-2]
        check('a draft named in Cyrillic gets a draft id and keeps its name',
              pid.startswith('draft') and load(pid)['name'] == 'Роман', pid)
        r = c.post(f'/project/{pid}/generate', follow_redirects=True)
        check('  and typesets', r.status_code == 200 and load(pid).get('last_pdf'), r.status_code)

        print('\n[browser: the real page]')
        browser = find_browser()
        if not browser:
            print('SKIP  no Chrome/Edge found (set TS_BROWSER to a Chromium browser)')
        else:
            from werkzeug.serving import make_server
            before = set(os.listdir(A.PROJECT_DIR))
            server = make_server('127.0.0.1', 0, A.app, threaded=True)
            threading.Thread(target=server.serve_forever, daemon=True).start()
            try:
                # real time: as in test_manuscript_editor.py
                proc = subprocess.Popen([browser, '--headless=new', '--disable-gpu',
                                         '--no-first-run', '--remote-debugging-port=0',
                                         '--user-data-dir=' + os.path.join(tmp, 'browser'),
                                         f'http://127.0.0.1:{server.server_port}/_fr_harness'],
                                        stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
                try:
                    for _ in range(240):
                        if RESULT:
                            break
                        time.sleep(0.5)
                finally:
                    proc.kill()
                    proc.wait()
            finally:
                server.shutdown()
            R = RESULT
            check('the harness ran to the end', R and not R.get('error'), R.get('error'))
            check('the page preview shows pages', R.get('pv_shown') is True, R.get('pv_cap'))
            check('the ebook preview shows the text', R.get('ev_shown') and R.get('ev_has_text'))
            check('Set a book reaches its result', R.get('result_has_save') is True)
            new = [n[:-5] for n in set(os.listdir(A.PROJECT_DIR)) - before if n.endswith('.json')]
            check('Save as project makes the book', R.get('after_save') == '/projects'
                  and len(new) == 1 and 'Текст первой главы' in
                  A._project_manuscript_text(load(new[0]), report_import=False), (R, new))
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    return 0 if not fails else 1


if __name__ == '__main__':
    code = main()
    print('\n' + ('ALL PASS' if not fails else 'FAILED: ' + ', '.join(fails)))
    sys.exit(code)
